"""Exports : CSV, Excel multi-feuilles et JSON complet.

L'export JSON inclut la charge brute (``raw``) de chaque enregistrement : c'est
la sortie a utiliser pour reexploiter la totalite des donnees collectees ailleurs.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Award, Buyer, SourceRun, Supplier, Tender, WatchProfile
from app.queries import (
    AwardFilters,
    TenderFilters,
    award_query,
    buyer_ranking,
    supplier_ranking,
    tender_query,
)

HEADER_FILL = PatternFill("solid", start_color="1D4ED8")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def _fmt(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if hasattr(value, "value"):  # enum
        return value.value
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


TENDER_HEADERS = [
    "Reference", "Objet", "Categorie", "Procedure", "Acheteur", "Service",
    "Lieu d'execution", "Lots", "Date de publication", "Date limite",
    "Seance d'ouverture", "Estimation", "Caution provisoire", "Devise",
    "Statut", "Score", "Mots-cles", "Source", "URL de l'avis",
    "Premiere detection", "Derniere detection",
]


def tender_rows(session: Session, filters: TenderFilters | None = None) -> list[list[Any]]:
    stmt = tender_query(filters or TenderFilters())
    return [
        [
            t.reference, t.object, _fmt(t.category), t.procedure_type, t.buyer_raw,
            t.buyer_service, t.execution_place, t.lots_count, _fmt(t.publication_date),
            _fmt(t.submission_deadline), _fmt(t.opening_session_date),
            t.estimated_amount, t.caution_provisoire, t.currency, _fmt(t.status),
            t.match_score, ", ".join(t.matched_keywords or []), t.source,
            t.consultation_url, _fmt(t.first_seen_at), _fmt(t.last_seen_at),
        ]
        for t in session.scalars(stmt).all()
    ]


AWARD_HEADERS = [
    "Reference", "Objet", "Lot", "Acheteur", "Attributaire", "Montant adjuge",
    "TTC/HT", "Estimation", "Devise", "Nombre de concurrents", "Date d'attribution",
    "Categorie", "Procedure", "Source", "URL", "Collecte le",
]


def award_rows(session: Session, filters: AwardFilters | None = None) -> list[list[Any]]:
    stmt = award_query(filters or AwardFilters())
    rows = []
    for a in session.scalars(stmt).all():
        tax = "TTC" if a.amount_includes_tax else ("HT" if a.amount_includes_tax is False else "")
        rows.append(
            [
                a.reference, a.object, a.lot_number, a.buyer_raw,
                (a.supplier.name if a.supplier else a.supplier_raw),
                a.awarded_amount, tax, a.estimated_amount, a.currency, a.bids_count,
                _fmt(a.award_date), _fmt(a.category), a.procedure_type, a.source,
                a.consultation_url, _fmt(a.scraped_at),
            ]
        )
    return rows


SUPPLIER_HEADERS = [
    "Attributaire", "Marches remportes", "Montant total", "Montant moyen",
    "Plus gros marche", "Acheteurs distincts", "Concurrents moyens",
    "Premier marche", "Dernier marche",
]


def supplier_rows(session: Session, limit: int = 5000) -> list[list[Any]]:
    return [
        [
            r.name, r.awards_count, r.total_amount, round(r.avg_amount, 2),
            r.max_amount, r.buyers_count,
            round(r.avg_bids, 1) if r.avg_bids is not None else None,
            _fmt(r.first_award), _fmt(r.last_award),
        ]
        for r in supplier_ranking(session, limit=limit)
    ]


BUYER_HEADERS = ["Acheteur", "Marches attribues", "Montant total", "Fournisseurs distincts"]


def buyer_rows(session: Session, limit: int = 5000) -> list[list[Any]]:
    return [
        [r.name, r.awards_count, r.total_amount, r.suppliers_count]
        for r in buyer_ranking(session, limit=limit)
    ]


def to_csv(headers: list[str], rows: Iterable[Iterable[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if cell is None else cell for cell in row])
    # BOM UTF-8 : Excel ouvre alors correctement les accents sans reglage.
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title[:31])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(["" if cell is None else cell for cell in row])
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        longest = max(
            [len(str(header))]
            + [len(str(row[index - 1])) for row in rows[:400] if row[index - 1] is not None]
            or [10]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 60)
    if rows:
        sheet.auto_filter.ref = sheet.dimensions


def to_xlsx(sheets: list[tuple[str, list[str], list[list[Any]]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in sheets:
        _write_sheet(workbook, title, headers, rows)
    if not workbook.sheetnames:
        workbook.create_sheet("Vide")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def full_workbook(session: Session) -> bytes:
    """Classeur complet : avis, attributions, attributaires, acheteurs, collectes."""
    runs = [
        [
            r.source, _fmt(r.started_at), _fmt(r.finished_at), r.status,
            r.items_fetched, r.items_new, r.items_updated, r.notifications_sent,
            r.estimated_cost_usd, r.error,
        ]
        for r in session.scalars(select(SourceRun).order_by(SourceRun.id.desc())).all()
    ]
    return to_xlsx(
        [
            ("Avis", TENDER_HEADERS, tender_rows(session)),
            ("Attributions", AWARD_HEADERS, award_rows(session)),
            ("Attributaires", SUPPLIER_HEADERS, supplier_rows(session)),
            ("Acheteurs", BUYER_HEADERS, buyer_rows(session)),
            (
                "Collectes",
                [
                    "Source", "Debut", "Fin", "Statut", "Recuperes", "Nouveaux",
                    "Mis a jour", "Notifications", "Cout USD", "Erreur",
                ],
                runs,
            ),
        ]
    )


def full_json(session: Session) -> bytes:
    """Export integral, charges brutes incluses : la sortie la plus complete."""

    def tender_dict(t: Tender) -> dict:
        return {
            "id": t.id, "source": t.source, "source_key": t.source_key,
            "reference": t.reference, "object": t.object, "category": _fmt(t.category),
            "procedure_type": t.procedure_type, "buyer": t.buyer_raw,
            "buyer_service": t.buyer_service, "execution_place": t.execution_place,
            "lots_count": t.lots_count, "publication_date": _fmt(t.publication_date),
            "submission_deadline": _fmt(t.submission_deadline),
            "opening_session_date": _fmt(t.opening_session_date),
            "estimated_amount": t.estimated_amount,
            "caution_provisoire": t.caution_provisoire, "currency": t.currency,
            "status": _fmt(t.status), "match_score": t.match_score,
            "matched_keywords": t.matched_keywords,
            "consultation_url": t.consultation_url,
            "first_seen_at": _fmt(t.first_seen_at), "last_seen_at": _fmt(t.last_seen_at),
            "raw": t.raw,
        }

    def award_dict(a: Award) -> dict:
        return {
            "id": a.id, "source": a.source, "source_key": a.source_key,
            "tender_id": a.tender_id, "reference": a.reference, "object": a.object,
            "lot_number": a.lot_number, "buyer": a.buyer_raw,
            "supplier": a.supplier.name if a.supplier else a.supplier_raw,
            "supplier_raw": a.supplier_raw, "awarded_amount": a.awarded_amount,
            "estimated_amount": a.estimated_amount, "currency": a.currency,
            "amount_includes_tax": a.amount_includes_tax, "bids_count": a.bids_count,
            "award_date": _fmt(a.award_date), "category": _fmt(a.category),
            "procedure_type": a.procedure_type, "consultation_url": a.consultation_url,
            "scraped_at": _fmt(a.scraped_at), "raw": a.raw,
        }

    payload = {
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "counts": {},
        "tenders": [tender_dict(t) for t in session.scalars(select(Tender)).all()],
        "awards": [award_dict(a) for a in session.scalars(select(Award)).all()],
        "suppliers": [
            {"id": s.id, "name": s.name, "normalized_name": s.normalized_name}
            for s in session.scalars(select(Supplier)).all()
        ],
        "buyers": [
            {"id": b.id, "name": b.name, "normalized_name": b.normalized_name,
             "acronym": b.acronym}
            for b in session.scalars(select(Buyer)).all()
        ],
        "watch_profiles": [
            {
                "id": p.id, "name": p.name, "active": p.active,
                "keywords": p.keywords, "excluded_keywords": p.excluded_keywords,
                "categories": p.categories, "min_amount": p.min_amount,
                "max_amount": p.max_amount, "min_score": p.min_score,
            }
            for p in session.scalars(select(WatchProfile)).all()
        ],
    }
    payload["counts"] = {key: len(payload[key]) for key in
                         ("tenders", "awards", "suppliers", "buyers", "watch_profiles")}
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")
